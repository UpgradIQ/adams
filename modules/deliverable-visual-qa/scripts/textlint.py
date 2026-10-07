import re, sys
import importlib.util, os
_spec = importlib.util.spec_from_file_location("hz", os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "humanize-writing", "scripts", "hzlint.py"))
_hz = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(_hz)
_G = _hz.load_profile()[1]
# Banned terms come from the active profile ("banned_terms": old brand names, words a project forbids); the dash follows the "dashes" group.
BANNED = list(_G.get("banned_terms", [])) + ([chr(0x2014)] if _G.get("dashes") else [])
AI_WORDS = re.compile(r'\b(delve|delving|leverage|leveraging|seamless|seamlessly|unlock|elevate|robust|game[- ]changer|cutting[- ]edge|tapestry|supercharge|empower)\b', re.I)
COUNT_EN = re.compile(r'\b(two|three|four|five|six|seven|eight|nine|ten)\s+(signs|habits|reasons|steps|tests|rules|questions|things|ways|mistakes|parts|stages|principles|tips|factors|lessons)\b', re.I)
COUNT_AR = re.compile(r'(ثلاث|تلات|أربع|خمس|ست|سبع|عشر)\S*\s+(علامات|عادات|أسباب|خطوات|قواعد|أسئلة|أخطاء|مراحل)')
FUNC = set('the a an of to and by for with on in as at from than that or but is are was were not'.split())
ICON = '[' + chr(0xF000) + '-' + chr(0xF8FF) + ']'  # built from code points so saving can never strip it

HITS = 0

def check(label, text):
    global HITS
    for b in BANNED:
        for m in re.finditer(re.escape(b), text):
            HITS += 1; print(label, 'BANNED', repr(text[max(0, m.start()-20):m.end()+20]))
    for rx, tag in ((AI_WORDS, 'AI-WORD'), (COUNT_EN, 'COUNT'), (COUNT_AR, 'COUNT')):
        for m in rx.finditer(text):
            HITS += 1; print(label, tag, m.group(0))

for f in sys.argv[1:]:
    if f.endswith('.pptx'):
        from pptx import Presentation  # only needed for pptx files
        p = Presentation(f)
        for i, s in enumerate(p.slides, 1):
            if not (s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()):
                print(f, i, 'NO SPEAKER NOTES')
            for sh in s.shapes:
                if not sh.has_text_frame:
                    continue
                t = sh.text_frame.text.strip()
                w = t.split()
                if 2 <= len(w) <= 9 and w[-1].lower().strip('.,') in FUNC:
                    print(f, i, 'TRUNCATED?', t)
                if re.search(ICON, t):
                    print(f, i, 'ICON GLYPH, use a vector icon')
                check(f'{f}#{i}', t)
    else:
        t = open(f, encoding='utf-8').read()
        check(f, t)
        if 'metadata' in f:
            ar = 'AR' in f
            head = '# وصف الفيديو' if ar else '# Video Description'
            tagh = '# التاجات' if ar else '# Tags'
            if head not in t or tagh not in t or '{{COURSE_URL}}' not in t:
                print(f, 'MISSING SECTION', head, tagh, '{{COURSE_URL}}')
                continue
            desc = t.split(head)[1].split('{{COURSE_URL}}')[0]
            tags = t.split(tagh)[1].strip()
            if not 3000 <= len(desc) <= 5000 or len(tags) > 500:
                print(f, 'LENGTH', len(desc), len(tags))
print('text lint hits:', HITS)
sys.exit(1 if HITS else 0)
